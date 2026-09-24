"""Xuất bảng nhập liệu ra Excel giữ nguyên công thức và ràng buộc, và nạp ngược tệp đó vào app.

Tệp xuất ra không phải ảnh chụp số liệu: tên tài khoản tra bằng VLOOKUP, cột Kiểm tra là công thức soát lỗi,
dòng Cộng dùng SUMIF, các ô nhập có danh sách chọn lấy từ danh mục của công ty, ô công thức bị khóa. Kế toán
gõ tiếp trong Excel rồi tải ngược lên, app đọc lại các cột giá trị nên không phụ thuộc Excel đã tính hay chưa.
"""
import base64
import io

from odoo import fields, models, _
from odoo.exceptions import UserError

from odoo.addons.lfood_importer.models.importer import to_date, to_number

SHEET, LOOKUP = 'NhatKy', 'DanhMuc'
# (nhãn, độ rộng, khóa ô)
COLUMNS = [('Ngày', 12, False), ('Số chứng từ', 16, False), ('Diễn giải', 38, False),
           ('TK Nợ', 10, False), ('Tên TK Nợ', 30, True), ('TK Có', 10, False), ('Tên TK Có', 30, True),
           ('Số tiền', 16, False), ('Đối tượng', 30, False), ('Khoản mục CP', 26, False), ('Kiểm tra', 46, True)]
FIRST = 4  # dòng dữ liệu đầu tiên (1-based) sau tiêu đề
BLANK = 50  # số dòng trống chừa sẵn để gõ tiếp


class LfoodEntrySheet(models.Model):
    _inherit = 'lfood.entry.sheet'

    file = fields.Binary('Tệp Excel đã xuất', readonly=True, copy=False, attachment=True)
    file_name = fields.Char(readonly=True, copy=False)
    import_file = fields.Binary('Tệp Excel nạp lên', copy=False, attachment=True)
    import_file_name = fields.Char(copy=False)

    # ------------------------------------------------------------------ xuất
    def _lookup_data(self):
        """Danh mục để Excel tra cứu: tài khoản, tài khoản phải ghi đối tượng, đối tượng, khoản mục chi phí."""
        self.ensure_one()
        accounts = self.env['lfood.account'].sudo().search([('allow_posting', '=', True)], order='code')
        partners = self.env['res.partner'].sudo().search([('is_company', '=', True)], order='name', limit=2000)
        items = self.env['lfood.cost.item'].sudo().search(
            ['|', ('company_id', '=', self.company_id.id), ('company_id', '=', False)], order='code')
        return accounts, partners, items

    def _write_lookup(self, book, formats):
        """Trang danh mục ẩn, đặt tên vùng để công thức và danh sách chọn dùng lại."""
        accounts, partners, items = self._lookup_data()
        sheet = book.add_worksheet(LOOKUP)
        sheet.write_row(0, 0, ['Tài khoản', 'Tên tài khoản', '', 'TK phải ghi đối tượng', '', 'Đối tượng', '', 'Khoản mục chi phí'],
                        formats['header'])
        for i, acc in enumerate(accounts, start=1):
            sheet.write_string(i, 0, acc.code)
            sheet.write_string(i, 1, acc.name or '')
        for i, acc in enumerate([a for a in accounts if a.track_partner], start=1):
            sheet.write_string(i, 3, acc.code)
        for i, p in enumerate(partners, start=1):
            sheet.write_string(i, 5, p.display_name)
        for i, it in enumerate(items, start=1):
            sheet.write_string(i, 7, it.display_name)
        book.define_name('TK', "=%s!$A$2:$A$%s" % (LOOKUP, max(len(accounts) + 1, 2)))
        book.define_name('TK_DoiTuong', "=%s!$D$2:$D$%s" % (LOOKUP, max(len([a for a in accounts if a.track_partner]) + 1, 2)))
        book.define_name('DoiTuong', "=%s!$F$2:$F$%s" % (LOOKUP, max(len(partners) + 1, 2)))
        book.define_name('KhoanMuc', "=%s!$H$2:$H$%s" % (LOOKUP, max(len(items) + 1, 2)))
        sheet.hide()
        sheet.protect()
        return len(accounts), len(partners), len(items)

    def _row_formulas(self, row):
        """Công thức của một dòng: tên tài khoản tra ngược, cột Kiểm tra soát đủ các ràng buộc."""
        r = row + 1  # Excel đếm từ 1
        name_of = lambda col: ('=IF(%(c)s%(r)s="","",IFERROR(VLOOKUP(%(c)s%(r)s,%(l)s!$A:$B,2,FALSE),'
                               '"không có tài khoản này"))' % {'c': col, 'r': r, 'l': LOOKUP})
        check = (
            '=IF(AND(A{r}="",D{r}="",F{r}=""),"",TRIM('
            'IF(A{r}="","Thiếu ngày. ","")'
            '&IF(B{r}="","Thiếu số chứng từ. ","")'
            '&IF(C{r}="","Thiếu diễn giải. ","")'
            '&IF(H{r}<=0,"Số tiền phải lớn hơn 0. ","")'
            '&IF(AND(D{r}="",F{r}=""),"Phải có ít nhất một tài khoản. ","")'
            '&IF(AND(D{r}<>"",D{r}=F{r}),"TK Nợ và TK Có trùng nhau. ","")'
            '&IF(AND(D{r}<>"",COUNTIF(TK,D{r})=0),"TK Nợ không có trong danh mục. ","")'
            '&IF(AND(F{r}<>"",COUNTIF(TK,F{r})=0),"TK Có không có trong danh mục. ","")'
            '&IF(AND(I{r}="",COUNTIF(TK_DoiTuong,D{r})+COUNTIF(TK_DoiTuong,F{r})>0),"Thiếu đối tượng. ","")'
            '))').format(r=r)
        return {4: name_of('D'), 6: name_of('F'), 10: check}

    def _build_workbook(self):
        """Trả về bytes của tệp .xlsx."""
        self.ensure_one()
        import xlsxwriter
        out = io.BytesIO()
        book = xlsxwriter.Workbook(out, {'in_memory': True, 'default_date_format': 'dd/mm/yyyy'})
        f = {
            'title': book.add_format({'bold': True, 'font_size': 13}),
            'header': book.add_format({'bold': True, 'bg_color': '#EEEEEE', 'border': 1, 'align': 'center',
                                       'text_wrap': True, 'locked': True}),
            'text': book.add_format({'border': 1, 'locked': False}),
            'date': book.add_format({'border': 1, 'locked': False, 'num_format': 'dd/mm/yyyy'}),
            'money': book.add_format({'border': 1, 'locked': False, 'num_format': '#,##0'}),
            'formula': book.add_format({'border': 1, 'locked': True, 'italic': True, 'font_color': '#555555'}),
            'total': book.add_format({'bold': True, 'border': 1, 'locked': True, 'num_format': '#,##0', 'bg_color': '#EEEEEE'}),
            'total_label': book.add_format({'bold': True, 'border': 1, 'locked': True, 'bg_color': '#EEEEEE'}),
            'bad': book.add_format({'bg_color': '#FFC7CE', 'font_color': '#9C0006'}),
        }
        sheet = book.add_worksheet(SHEET)
        self._write_lookup(book, f)  # sau trang chính thì Excel mới cho ẩn trang danh mục
        sheet.write(0, 0, '%s - %s' % (self.company_id.name, self.name), f['title'])
        sheet.write(1, 0, _('Chỉ gõ vào ô trắng. Cột Tên TK và Kiểm tra là công thức, đã khóa. '
                            'Gõ xong tải tệp này lên lại app ở nút Nạp từ Excel.'))
        for col, (label, width, _locked) in enumerate(COLUMNS):
            sheet.write(FIRST - 2, col, label, f['header'])
            sheet.set_column(col, col, width)
        sheet.freeze_panes(FIRST - 1, 0)

        lines = self.line_ids.sorted(lambda l: (l.sequence, l.id))
        last = FIRST - 1 + len(lines) + BLANK
        for i in range(len(lines) + BLANK):
            row = FIRST - 1 + i
            line = lines[i] if i < len(lines) else None
            sheet.write_datetime(row, 0, line.date, f['date']) if line and line.date else sheet.write_blank(row, 0, None, f['date'])
            sheet.write_string(row, 1, line.ref or '' if line else '', f['text'])
            sheet.write_string(row, 2, line.memo or '' if line else '', f['text'])
            sheet.write_string(row, 3, line.debit_account_id.code or '' if line else '', f['text'])
            sheet.write_string(row, 5, line.credit_account_id.code or '' if line else '', f['text'])
            sheet.write_number(row, 7, line.amount or 0, f['money']) if line else sheet.write_blank(row, 7, None, f['money'])
            sheet.write_string(row, 8, line.partner_id.display_name or '' if line else '', f['text'])
            sheet.write_string(row, 9, line.cost_item_id.display_name or '' if line else '', f['text'])
            for col, formula in self._row_formulas(row).items():
                sheet.write_formula(row, col, formula, f['formula'], '')

        total = last  # dòng cộng nằm ngay dưới vùng dữ liệu (0-based)
        first_excel, last_excel = FIRST, last
        sheet.write(total, 2, _('Cộng'), f['total_label'])
        sheet.write_formula(total, 3, '=SUMIF(D%s:D%s,"<>",$H$%s:$H$%s)' % (first_excel, last_excel, first_excel, last_excel), f['total'])
        sheet.write_formula(total, 5, '=SUMIF(F%s:F%s,"<>",$H$%s:$H$%s)' % (first_excel, last_excel, first_excel, last_excel), f['total'])
        sheet.write_formula(total, 7, '=SUM(H%s:H%s)' % (first_excel, last_excel), f['total'])
        sheet.write_formula(total, 10, '=IF(D%(t)s=F%(t)s,"Nợ bằng Có","Lệch "&TEXT(D%(t)s-F%(t)s,"#,##0"))'
                            % {'t': total + 1}, f['total_label'])

        # ràng buộc nhập: ngày, danh sách tài khoản, số tiền dương, đối tượng, khoản mục
        rng = (FIRST - 1, 0, last - 1, 0)
        sheet.data_validation(*rng, {'validate': 'date', 'criteria': 'between',
                                     'minimum': fields.Date.to_date('2000-01-01'), 'maximum': fields.Date.to_date('2099-12-31'),
                                     'error_title': 'Ngày không hợp lệ', 'error_message': 'Nhập ngày dạng dd/mm/yyyy.'})
        for col in (3, 5):
            sheet.data_validation(FIRST - 1, col, last - 1, col,
                                  {'validate': 'list', 'source': '=TK', 'error_title': 'Tài khoản không có',
                                   'error_message': 'Chọn tài khoản chi tiết trong danh mục của công ty.'})
        sheet.data_validation(FIRST - 1, 7, last - 1, 7,
                              {'validate': 'decimal', 'criteria': '>', 'value': 0, 'error_title': 'Số tiền',
                               'error_message': 'Số tiền phải lớn hơn 0.'})
        sheet.data_validation(FIRST - 1, 8, last - 1, 8, {'validate': 'list', 'source': '=DoiTuong', 'error_type': 'warning',
                                                          'error_title': 'Đối tượng', 'error_message': 'Nên chọn trong danh mục đối tượng.'})
        sheet.data_validation(FIRST - 1, 9, last - 1, 9, {'validate': 'list', 'source': '=KhoanMuc', 'error_type': 'warning',
                                                          'error_title': 'Khoản mục', 'error_message': 'Nên chọn trong danh mục khoản mục chi phí.'})
        sheet.conditional_format(FIRST - 1, 10, last - 1, 10,
                                 {'type': 'cell', 'criteria': '!=', 'value': '""', 'format': f['bad']})
        sheet.protect('', {'select_locked_cells': True, 'select_unlocked_cells': True, 'sort': True, 'autofilter': True})
        book.close()
        return out.getvalue()

    def action_export_excel(self):
        self.ensure_one()
        data = self._build_workbook()
        self.sudo().write({'file': base64.b64encode(data),
                           'file_name': '%s.xlsx' % (self.name or 'bang-nhap-lieu').replace('/', '-')})
        self.env['lfood.audit.log']._record_event(
            'export', model=self._name, res_id=self.id, res_name=self.name, company_id=self.company_id.id,
            summary=_('Xuất bảng nhập liệu %s ra Excel: %s dòng') % (self.name, len(self.line_ids)))
        return {'type': 'ir.actions.act_url', 'target': 'self',
                'url': '/web/content/%s/%s/file/%s?download=true' % (self._name, self.id, self.file_name)}

    # ------------------------------------------------------------------ nạp ngược
    def action_import_excel(self):
        """Đọc lại tệp đã xuất (hoặc tệp cùng mẫu) và thay toàn bộ các dòng của bảng."""
        self.ensure_one()
        if self.state != 'draft':
            raise UserError(_('Bảng đã ghi sổ không nạp lại được.'))
        if not self.import_file:
            raise UserError(_('Chưa chọn tệp Excel.'))
        import openpyxl
        raw = base64.b64decode(self.import_file)
        if raw[:2] != b'PK':
            raise UserError(_('Tệp phải là .xlsx xuất từ app hoặc cùng mẫu cột.'))
        book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = book[SHEET] if SHEET in book.sheetnames else book.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(min_row=FIRST, values_only=True)]

        Account = self.env['lfood.account'].sudo()
        accounts = {a.code: a.id for a in Account.search([('allow_posting', '=', True)])}
        partners = {p.display_name: p.id for p in self.env['res.partner'].sudo().search([])}
        items = {i.display_name: i.id for i in self.env['lfood.cost.item'].sudo().search(
            ['|', ('company_id', '=', self.company_id.id), ('company_id', '=', False)])}

        def code_of(v):
            code = str(v).strip() if v not in (None, '') else ''
            return code.split(' ')[0] if code else ''

        vals, problems = [], []
        for n, row in enumerate(rows, start=FIRST):
            row = list(row) + [None] * (len(COLUMNS) - len(row))
            date, ref, memo, debit, _dn, credit, _cn, amount, partner, item = row[:10]
            if not any(v not in (None, '') for v in (date, ref, memo, debit, credit, amount)):
                continue
            if str(memo or '').strip() == 'Cộng':
                continue
            debit, credit = code_of(debit), code_of(credit)
            for code in (debit, credit):
                if code and code not in accounts:
                    problems.append(_('Dòng %s: không có tài khoản chi tiết %s') % (n, code))
            partner_name = str(partner).strip() if partner else ''
            item_name = str(item).strip() if item else ''
            if partner_name and partner_name not in partners:
                problems.append(_('Dòng %s: không có đối tượng "%s"') % (n, partner_name))
            if item_name and item_name not in items:
                problems.append(_('Dòng %s: không có khoản mục "%s"') % (n, item_name))
            vals.append({'sequence': (n - FIRST + 1) * 10, 'date': to_date(date), 'ref': str(ref).strip() if ref else False,
                         'memo': str(memo).strip() if memo else False,
                         'debit_account_id': accounts.get(debit, False), 'credit_account_id': accounts.get(credit, False),
                         'amount': to_number(amount), 'partner_id': partners.get(partner_name, False),
                         'cost_item_id': items.get(item_name, False)})
        if problems:
            raise UserError(_('Tệp còn %s chỗ chưa khớp danh mục:\n%s') % (len(problems), '\n'.join(problems[:20])))
        if not vals:
            raise UserError(_('Tệp không có dòng dữ liệu nào.'))
        self.line_ids.unlink()
        self.write({'line_ids': [(0, 0, v) for v in vals], 'import_file': False, 'import_file_name': False})
        self.env['lfood.audit.log']._record_event(
            'write', model=self._name, res_id=self.id, res_name=self.name, company_id=self.company_id.id,
            summary=_('Nạp bảng nhập liệu %s từ Excel: %s dòng') % (self.name, len(vals)))
        return self._check_lines()
